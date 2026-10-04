"""
Generate SQL to align database 100% 1:1 with Authoritative Scorecard Workbooks
"""

import openpyxl
import os
import re

base_27_sep = r'docs/PAKET_FINAL_SCORECARD_27_SEPTEMBER_DAN_BASELINE (1)/PAKET_FINAL_SCORECARD_27_SEPTEMBER_DAN_BASELINE/01_SCORECARD_FINAL_27_SEPTEMBER_2026'

sc_files = {
    1: 'Scorecard_P01_DEKRANASDA_KEPRI_27_Sep_2026.xlsx',
    2: 'Scorecard_P02_BNNP_MASA_IMPLEMENTASI_AWAL_27_Sep_2026.xlsx',
    3: 'Scorecard_P03_PEMKOT_TANJUNGPINANG_FINAL_27_Sep_2026.xlsx',
    4: 'Scorecard_P04_BAPPERIDA_BINTAN_27_Sep_2026.xlsx',
    5: 'Scorecard_P05_STAIN_SAR_KEPRI_27_Sep_2026.xlsx',
    6: 'Scorecard_P06_UMRAH_27_Sep_2026.xlsx',
    7: 'Scorecard_P07_STAI_ANAMBAS_27_Sep_2026.xlsx',
    8: 'Scorecard_P08_POLIBATAM_27_Sep_2026.xlsx',
    9: 'Scorecard_P09_STAI_NATUNA_27_Sep_2026.xlsx',
    10: 'Scorecard_P10_STISIP_BATAM_27_Sep_2026.xlsx',
}

weights_v21 = {'I1': 10, 'I2': 15, 'I3': 15, 'I4': 20, 'I5': 20, 'I6': 10, 'I7': 10}

def esc(s):
    if s is None:
        return 'NULL'
    return "'" + str(s).replace('\\', '\\\\').replace("'", "''") + "'"

sql_lines = []
sql_lines.append("-- ========================================================")
sql_lines.append("-- SQL SCRIPT TO ALIGN DATABASE 100% 1:1 WITH SCORECARDS")
sql_lines.append("-- ========================================================")
sql_lines.append("USE mitra_kinerja;\n")

# Canonical deskripsi extracted from authoritative 27 Sep workbooks
canonical_desc = {}
sample_wb = openpyxl.load_workbook(os.path.join(base_27_sep, sc_files[1]), data_only=True)
sample_pen = sample_wb['PENILAIAN']
for r in range(13, 20):
    c1 = str(sample_pen.cell(r, 1).value or '').strip()
    m_code = re.match(r'^(I[1-7])', c1)
    if m_code:
        ind = m_code.group(1)
        title = c1.split('. ', 1)[-1].strip()
        apa_dinilai = str(sample_pen.cell(r, 3).value or '').strip()
        canonical_desc[ind] = f"{title}\nCara periksa: {apa_dinilai}"

# 1. Update deskripsi for all 18 partners to standard authoritative text
sql_lines.append("-- 1. Update deskripsi for all 18 partners to authoritative Result-Chain standard")
for ind, desc in canonical_desc.items():
    sql_lines.append(f"UPDATE indikator_skor SET deskripsi = {esc(desc)} WHERE kode_indikator = '{ind}';")
sql_lines.append("")

# 2. Update MIDs 1 to 10 from 27 September workbooks
sql_lines.append("-- 2. Update MIDs 1-10 indicators and metadata from 27 September workbooks")
for mid, fname in sorted(sc_files.items()):
    fpath = os.path.join(base_27_sep, fname)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    pen = wb['PENILAIAN']
    rek = wb['REKOMENDASI']

    wb_status = str(pen.cell(27, 2).value or '').strip()
    wb_posisi = str(rek.cell(9, 2).value or '').strip()
    wb_rekom = str(rek.cell(17, 2).value or '').strip()
    wb_temuan = str(rek.cell(16, 2).value or '').strip()
    wb_tl = str(rek.cell(19, 2).value or '').strip()
    expected_full_rekom = f"{wb_rekom}\n\nTemuan Utama:\n{wb_temuan}\n\nRencana Tindak Lanjut:\n{wb_tl}".strip()

    sql_lines.append(f"-- Partner MID {mid}: {fname}")
    sql_lines.append(f"UPDATE mitra_kinerja SET status_scorecard = {esc(wb_status)}, posisi_portofolio = {esc(wb_posisi)}, rekomendasi = {esc(expected_full_rekom)} WHERE id = {mid};")

    for r in range(13, 20):
        c1 = str(pen.cell(r, 1).value or '').strip()
        m_code = re.match(r'^(I[1-7])', c1)
        if not m_code: continue
        ind = m_code.group(1)

        wb_bobot = weights_v21[ind]
        wb_bl = str(pen.cell(r, 4).value or '').strip()
        wb_ini = str(pen.cell(r, 5).value or '').strip()
        wb_st = str(pen.cell(r, 6).value or '').strip()
        wb_bukti = str(pen.cell(r, 7).value or '').strip()
        wb_skor = pen.cell(r, 8).value
        wb_nilai = pen.cell(r, 9).value
        wb_alasan = str(pen.cell(r, 10).value or '').strip()
        wb_tl = str(pen.cell(r, 11).value or '').strip()

        wb_skor_sql = "NULL" if wb_skor is None or str(wb_skor).strip() == '' else str(int(wb_skor))
        wb_nilai_sql = "NULL" if wb_nilai is None or str(wb_nilai).strip() == '' else f"{float(wb_nilai):.2f}"

        sql_ind = f"""UPDATE indikator_skor SET 
    bobot = {wb_bobot},
    kondisi_baseline = {esc(wb_bl)},
    kondisi_saat_ini = {esc(wb_ini)},
    status_pemeriksaan = {esc(wb_st)},
    temuan_bukti = {esc(wb_bukti)},
    skor = {wb_skor_sql},
    nilai = {wb_nilai_sql},
    alasan_skor = {esc(wb_alasan)},
    catatan_tindak_lanjut = {esc(wb_tl)},
    referensi_baseline = {esc(wb_bl)}
WHERE mitra_id = {mid} AND kode_indikator = '{ind}';"""
        sql_lines.append(sql_ind)
    sql_lines.append("")

output_sql = "\n".join(sql_lines)
with open('database/align_scorecard_100_parity.sql', 'w', encoding='utf-8') as f:
    f.write(output_sql)

print(f"Alignment SQL generated: database/align_scorecard_100_parity.sql ({len(output_sql)} bytes)")
