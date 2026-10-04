"""
Generate SQL for optional ingestion of legacy V1.1 scorecards (C02, P05 Mumtaz, C03, C04) into V2.1 schema
"""

import openpyxl
import os

base_final = r'docs/Paket_Scorecard_dan_Baseline_Final/Paket_Scorecard_dan_Baseline_Final'

legacy_map = {
    11: ('02_Portofolio_Pengayaan', 'C02_Scorecard_Politeknik_Bintan_Cakrawala.xlsx', 'SCORECARD PBC', 'BELUM LENGKAP'),
    12: ('01_Pilot_Utama', 'P05_Scorecard_STIT_Mumtaz_Karimun.xlsx', 'SCORECARD MUMTAZ', 'BELUM LENGKAP'),
    13: ('02_Portofolio_Pengayaan', 'C03_Scorecard_Universitas_Ibnu_Sina.xlsx', 'SCORECARD UIS', 'SIAP DIVALIDASI'),
    14: ('02_Portofolio_Pengayaan', 'C04_Scorecard_UNRIKA.xlsx', 'SCORECARD UNRIKA', 'BELUM LENGKAP'),
}

weights_v21 = {'I1': 10, 'I2': 15, 'I3': 15, 'I4': 20, 'I5': 20, 'I6': 10, 'I7': 10}

def esc(s):
    if s is None:
        return 'NULL'
    return "'" + str(s).replace('\\', '\\\\').replace("'", "''") + "'"

sql_lines = []
sql_lines.append("-- ==========================================================================")
sql_lines.append("-- OPTIONAL SQL SCRIPT: INGEST LEGACY V1.1 SCORECARDS (MIDs 11-14) INTO V2.1")
sql_lines.append("-- Note: Ingests I1-I6 and maps weights to V2.1 Result-Chain standards.")
sql_lines.append("-- I7 remains 'BELUM DITELAAH' pending formal V2.1 Result-Chain evaluation.")
sql_lines.append("-- ==========================================================================")
sql_lines.append("USE mitra_kinerja;\n")

for mid, (group, fname, sheet_name, default_status) in legacy_map.items():
    fpath = os.path.join(base_final, group, fname)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb[sheet_name]

    wb_status = str(ws.cell(21, 8).value or default_status).strip()
    sql_lines.append(f"-- Partner MID {mid}: {fname}")
    sql_lines.append(f"UPDATE mitra_kinerja SET status_scorecard = {esc(wb_status)} WHERE id = {mid};")

    for r in range(13, 19):
        ind = str(ws.cell(r, 1).value or '').strip()
        wb_bl = str(ws.cell(r, 4).value or '').strip()
        wb_st_raw = str(ws.cell(r, 5).value or '').strip()
        wb_bukti = str(ws.cell(r, 6).value or '').strip()
        wb_skor = ws.cell(r, 7).value
        wb_alasan = str(ws.cell(r, 8).value or '').strip()

        bobot = weights_v21[ind]

        if 'MEMADAI' in wb_st_raw and 'BELUM' not in wb_st_raw:
            st_db = 'DAPAT DINILAI'
        elif 'BELUM MEMADAI' in wb_st_raw:
            st_db = 'BUKTI BELUM MEMADAI'
        else:
            st_db = 'BELUM DITELAAH'

        if st_db == 'DAPAT DINILAI' and wb_skor is not None and str(wb_skor).strip() != '':
            skor_int = int(float(wb_skor))
            nilai_val = round((skor_int / 4.0) * bobot, 2)
            skor_sql = str(skor_int)
            nilai_sql = f"{nilai_val:.2f}"
        else:
            skor_sql = "NULL"
            nilai_sql = "NULL"

        sql_ind = f"""UPDATE indikator_skor SET 
    bobot = {bobot},
    kondisi_baseline = {esc(wb_bl)},
    kondisi_saat_ini = {esc(wb_bukti)},
    status_pemeriksaan = {esc(st_db)},
    temuan_bukti = {esc(wb_bukti)},
    skor = {skor_sql},
    nilai = {nilai_sql},
    alasan_skor = {esc(wb_alasan)},
    referensi_baseline = {esc(wb_bl)}
WHERE mitra_id = {mid} AND kode_indikator = '{ind}';"""
        sql_lines.append(sql_ind)
    sql_lines.append("")

output_sql = "\n".join(sql_lines)
with open('database/align_legacy_v11_scorecards.sql', 'w', encoding='utf-8') as f:
    f.write(output_sql)

print(f"Legacy V1.1 SQL generated: database/align_legacy_v11_scorecards.sql ({len(output_sql)} bytes)")
